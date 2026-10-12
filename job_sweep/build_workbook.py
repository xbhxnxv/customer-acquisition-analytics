"""Build Best_Fit_Jobs_By_Deadline.xlsx and SUMMARY.md from progress.jsonl (Section 6 of the brief)."""
import csv
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from score import fit_score, role_type

HERE = Path(__file__).parent
TODAY = date.today()
MIN_SCORE = 55
MAX_PER_COMPANY = 3

RED, AMBER, GREEN, GREY = ("FFC7CE", "FFEB9C", "C6EFCE", "D9D9D9")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
WRAP = Alignment(wrap_text=True, vertical="top")

MAIN_COLS = [
    ("Priority", 8), ("Deadline", 12), ("Days Left", 9), ("Company", 22), ("Region", 8),
    ("Role Title", 40), ("Role Type", 16), ("Location", 20), ("Start Date", 14), ("Salary", 18),
    ("Fit Score", 8), ("Why It Fits", 50), ("Key Requirements", 45), ("Application Steps", 40),
    ("Sponsorship", 16), ("Already Applied to This Company?", 12), ("Apply Link", 45),
    ("Deadline Source", 22), ("Original Deadline Text", 30), ("Posted Date", 12), ("Notes", 45),
    ("Verified On", 12),
]


def load_progress():
    """Latest record per company number."""
    recs = {}
    p = HERE / "progress.jsonl"
    if p.exists():
        for line in p.read_text().splitlines():
            if line.strip():
                r = json.loads(line)
                recs[r["n"]] = r
    return recs


def load_companies():
    with open(HERE / "companies.csv") as f:
        return list(csv.DictReader(f))


def parse_deadline(d):
    if d in (None, "", "Rolling"):
        return None
    return date.fromisoformat(d)


def collect(recs, companies):
    applied = {int(c["#"]): c["Applied"] == "YES" for c in companies}
    main, closed, opening = [], [], []
    for n, rec in sorted(recs.items()):
        kept = []
        for r in rec.get("roles", []):
            score = fit_score(r)
            dl = parse_deadline(r["deadline"])
            if dl and dl < TODAY:
                closed.append([rec["company"], r["title"], r["deadline"], "Deadline passed", r["link"]])
                continue
            if score < MIN_SCORE:
                closed.append([rec["company"], r["title"], r["deadline"], f"Fit score {score} < {MIN_SCORE}", r["link"]])
                continue
            kept.append((score, dl, r))
        kept.sort(key=lambda t: -t[0])
        for score, dl, r in kept[:MAX_PER_COMPANY]:
            main.append(dict(r, company=rec["company"], region=rec["region"], score=score, dl=dl,
                             applied="Yes" if applied.get(n) else "No"))
        for score, dl, r in kept[MAX_PER_COMPANY:]:
            closed.append([rec["company"], r["title"], r["deadline"], "Over 3-per-company cap (lower fit)", r["link"]])
        for c in rec.get("closed", []):
            closed.append([rec["company"], c["role"], c.get("deadline", ""), c["reason"], c.get("link", "")])
        for o in rec.get("opening", []):
            opening.append([rec["company"], o["role"], o.get("opens", "Not stated"),
                            o.get("exp_deadline", "Not stated"), o.get("link", ""), o.get("notes", "")])
    dated = sorted([m for m in main if m["dl"]], key=lambda m: (m["dl"], -m["score"]))
    rolling = [m for m in main if not m["dl"]]
    # Fit Score desc, then newest Posted Date first
    rolling.sort(key=lambda m: (-m["score"], -(date.fromisoformat(m["posted"]).toordinal() if m.get("posted") else 0)))
    return dated + rolling, closed, opening


def style_header(ws, widths):
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def simple_sheet(wb, title, headers, rows, widths, link_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
        if link_col is not None and r[link_col]:
            c = ws.cell(row=ws.max_row, column=link_col + 1)
            if str(r[link_col]).startswith("http"):
                c.hyperlink = r[link_col]
                c.font = Font(color="0563C1", underline="single")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    style_header(ws, widths)
    return ws


def build():
    companies = load_companies()
    recs = load_progress()
    main, closed, opening = collect(recs, companies)

    wb = Workbook()
    ws = wb.active
    ws.title = "Apply By Deadline"
    ws.append([c for c, _ in MAIN_COLS])
    for i, m in enumerate(main, 1):
        days = (m["dl"] - TODAY).days if m["dl"] else None
        ws.append([
            i, m["dl"] if m["dl"] else "Rolling", days, m["company"], m["region"], m["title"],
            role_type(m), m["loc"], m.get("start", "Not stated"), m.get("salary", "Not stated"),
            m["score"], m.get("why", ""), m.get("reqs", "Not stated"), m.get("steps", "Not stated"),
            m.get("sponsor", "Not stated"), m["applied"], m["link"], m["dl_source"], m["dl_text"],
            m.get("posted", ""), m.get("notes", ""), TODAY,
        ])
        row = ws.max_row
        ws.cell(row=row, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=22).number_format = "yyyy-mm-dd"
        if days is None:
            colour = GREY
        elif days <= 7:
            colour = RED
        elif days <= 21:
            colour = AMBER
        else:
            colour = GREEN
        for col in (2, 3):
            ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=colour)
        link = ws.cell(row=row, column=17)
        link.hyperlink = m["link"]
        link.font = Font(color="0563C1", underline="single")
        for c in ws[row]:
            c.alignment = WRAP
    style_header(ws, [w for _, w in MAIN_COLS])

    simple_sheet(wb, "Opening Soon", ["Company", "Role", "Opens On", "Expected Deadline", "Link", "Notes"],
                 opening, [24, 45, 16, 18, 50, 50], link_col=4)
    simple_sheet(wb, "Closed - Excluded", ["Company", "Role", "Deadline", "Reason", "Link"],
                 closed, [24, 50, 14, 45, 55], link_col=4)

    cov_rows, manual = [], []
    listed = {}
    for m in main:
        listed[(m["company"], m["region"])] = listed.get((m["company"], m["region"]), 0) + 1
    for c in companies:
        n = int(c["#"])
        rec = recs.get(n)
        if rec is None:
            cov_rows.append([n, c["Company"], c["Region"], "NOT YET CHECKED", 0, "", c["Careers URL"], ""])
            manual.append(f'{c["Company"]} ({c["Region"]}) — not yet checked')
            continue
        k = listed.get((rec["company"], rec["region"]), 0)
        status = rec["status"]
        if status == "Roles found" and k == 0:
            status = "No suitable role open"
        cov_rows.append([n, rec["company"], rec["region"], status, k, rec["method"], rec["careers_url"], rec.get("notes", "")])
        if status.startswith(("Blocked", "Careers site not found")) or rec.get("manual"):
            manual.append(f'{rec["company"]} ({rec["region"]}) — {rec.get("manual") or status}')
    simple_sheet(wb, "Coverage", ["#", "Company", "Region", "Status", "No. of roles listed", "Method used",
                                  "Careers URL actually used", "Notes"],
                 cov_rows, [5, 30, 8, 24, 10, 22, 50, 60], link_col=6)

    within = lambda d: sum(1 for m in main if m["dl"] and (m["dl"] - TODAY).days <= d)
    readme = wb.create_sheet("Read Me")
    lines = [
        ("Run date", TODAY.isoformat()),
        ("Companies in scope", len(companies)),
        ("Companies checked", len(recs)),
        ("Roles listed (main sheet)", len(main)),
        ("  with a dated deadline", sum(1 for m in main if m["dl"])),
        ("  rolling / no deadline", sum(1 for m in main if not m["dl"])),
        ("Closing in next 7 days", within(7)),
        ("Closing in next 14 days", within(14)),
        ("Closing in next 30 days", within(30)),
        ("", ""),
        ("SETTINGS", ""),
        ("START_DATE_FILTER", "2026 and 2027 starts included; 2027 starts flagged '2027 START' in Notes"),
        ("ROLLING_ROLES", "Included, listed after dated roles (by Fit Score, then Posted Date)"),
        ("EXPIRED_ROLES", "Excluded from main list; logged in 'Closed / Excluded'"),
        ("ALREADY_APPLIED_COMPANIES", "Searched; only roles different from the one already applied for are listed"),
        ("MAX_ROLES_PER_COMPANY", MAX_PER_COMPANY),
        ("MIN FIT SCORE", MIN_SCORE),
        ("REGIONS", "UK (anywhere) and India metros; remote open to UK/India"),
        ("SPONSORSHIP", "Not filtered; recorded as stated"),
        ("", ""),
        ("METHOD / ASSUMPTIONS", ""),
    ]
    for a in json.loads((HERE / "assumptions.json").read_text()) if (HERE / "assumptions.json").exists() else []:
        lines.append(("", a))
    reminders = applied_reminders(recs)
    lines += [("", ""), ("ALREADY-APPLIED: LIVE DEADLINES", "Same role you applied for — confirm your application covers this cycle")]
    lines += [(r[0], r[1]) for r in reminders]
    lines += [("", ""), ("CHECK MANUALLY", "")] + [("", m) for m in manual]
    for k, v in lines:
        readme.append([k, v])
    readme.column_dimensions["A"].width = 30
    readme.column_dimensions["B"].width = 120
    for row in readme.iter_rows():
        for c in row:
            c.alignment = WRAP
        if row[0].value in ("SETTINGS", "METHOD / ASSUMPTIONS", "CHECK MANUALLY", "ALREADY-APPLIED: LIVE DEADLINES"):
            row[0].font = Font(bold=True)

    out = HERE / "Best_Fit_Jobs_By_Deadline.xlsx"
    wb.save(out)
    write_summary(main, manual, reminders)
    return out, main, recs


def applied_reminders(recs):
    """(label, detail) for already-applied companies whose same role has a live deadline, nearest first."""
    out = []
    for rec in recs.values():
        r = rec.get("reminder")
        if r and (r["deadline"] == "Rolling" or date.fromisoformat(r["deadline"]) >= TODAY):
            out.append((r["deadline"], f'{rec["company"]}: {r["role"]} — {r["link"]}'))
    return sorted(out, key=lambda t: (t[0] == "Rolling", t[0]))


def write_summary(main, manual, reminders):
    def link(m):
        return f'[{m["title"]}]({m["link"]})'
    dated = [m for m in main if m["dl"]][:10]
    top = sorted(main, key=lambda m: -m["score"])[:10]
    out = [f"# Best-fit job sweep — {TODAY.isoformat()}", "",
           f"{len(main)} roles listed in `Best_Fit_Jobs_By_Deadline.xlsx`.", "",
           "## 10 nearest deadlines", "", "| Deadline | Days | Company | Role | Fit |", "|---|---|---|---|---|"]
    out += [f'| {m["dl"]} | {(m["dl"] - TODAY).days} | {m["company"]} ({m["region"]}) | {link(m)} | {m["score"]} |' for m in dated]
    out += ["", "## Top 10 by Fit Score", "", "| Fit | Deadline | Company | Role |", "|---|---|---|---|"]
    out += [f'| {m["score"]} | {m["dl"] or "Rolling"} | {m["company"]} ({m["region"]}) | {link(m)} |' for m in top]
    out += ["", "## Already-applied companies with a live deadline for the same role",
            "", "Confirm your earlier application covers this (2027) cycle:", ""]
    out += [f"- **{d}** — {t}" for d, t in reminders]
    out += ["", "## Check manually", ""] + [f"- {x}" for x in manual]
    (HERE / "SUMMARY.md").write_text("\n".join(out) + "\n")


if __name__ == "__main__":
    out, main, recs = build()
    print(f"{out.name}: {len(recs)} companies checked, {len(main)} roles listed")
